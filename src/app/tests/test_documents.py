"""Regression tests for the proposal PDF generator (src/app/server/documents.py).

Covers two bugs fixed together:
1. The table only ever showed the customer's raw_description — never the actual
   catalog product a line was matched/priced against. `_matched_product_description`
   is the pure lookup that fixes that; it's tested directly here.
2. Long cells (a product description, a verbose raw_description) used to overflow
   into neighboring columns because they were plain strings instead of wrapped
   Paragraph flowables. That's covered by a smoke test that renders a bid with a
   deliberately long description and asserts it doesn't blow up — the actual
   word-wrap behavior is ReportLab's own, layout-tested visually, but this at
   least guards against a Paragraph/column-count regression.
"""

from __future__ import annotations

from io import BytesIO

import pytest

from server.documents import _matched_product_description, render_customer_submission_excel, render_proposal_pdf


def _candidate(item_code: str, **overrides: object) -> dict[str, object]:
    base: dict[str, object] = {
        "item_code": item_code,
        "description_long": "Kimberly-Clark Kimtech M nitrile exam gloves, powder-free, textured fingertips, 100/BX",
        "manufacturer_name": "Kimberly-Clark",
        "score": 0.89,
        "source": "purchase_history",
        "rationale": "82% catalog match, boosted by purchase/bid history to 89%.",
    }
    base.update(overrides)
    return base


def _line(line_number: int, **overrides: object) -> dict[str, object]:
    base: dict[str, object] = {
        "line_id": f"line-{line_number}",
        "line_number": line_number,
        "raw_description": "Nitrile exam gloves, powder free, size medium, box of 100ct",
        "qty": 200,
        "uom": "BX",
        "candidates": [_candidate("SKU-100005")],
        "selected_item_code": "SKU-100005",
        "proposed_price": 8.17,
        "pricing_basis": "reviewer-selected",
        "review_action": "accepted",
    }
    base.update(overrides)
    return base


def test_matched_product_description_uses_selected_candidate():
    line = _line(1)
    assert _matched_product_description(line) == (
        "Kimberly-Clark — Kimberly-Clark Kimtech M nitrile exam gloves, "
        "powder-free, textured fingertips, 100/BX"
    )


def test_matched_product_description_without_manufacturer():
    line = _line(1, candidates=[_candidate("SKU-100005", manufacturer_name=None)])
    assert _matched_product_description(line) == (
        "Kimberly-Clark Kimtech M nitrile exam gloves, powder-free, textured fingertips, 100/BX"
    )


def test_matched_product_description_falls_back_when_unmatched():
    line = _line(1, selected_item_code=None, candidates=[])
    assert _matched_product_description(line) == "—"


def test_matched_product_description_falls_back_when_selection_not_in_candidates():
    """Defensive: selected_item_code should always be one of the line's own
    candidates, but don't crash if data drifts (e.g. a stale manual-match)."""
    line = _line(1, selected_item_code="SKU-999999")
    assert _matched_product_description(line) == "—"


def test_render_proposal_pdf_handles_long_cells_without_raising():
    """The real bug report: a long raw_description used to visually overflow
    into the qty/UOM columns because cells were plain strings. Wrapping them in
    Paragraph flowables is what fixes that (ReportLab wraps a Paragraph within
    its column); this just guards against a gross regression (wrong arg types,
    mismatched column count vs. header, etc.) — not exact pixel layout."""
    pytest.importorskip("reportlab")
    bid = {
        "bid_id": "bid-test",
        "customer_name": "Test Customer",
        "solicitation_ref": "",
        "line_items": [
            _line(
                1,
                raw_description="A very long requested-product description that is deliberately much "
                "longer than the column width so it would have overflowed under the old plain-string "
                "table cells instead of wrapping to a second line",
                candidates=[_candidate(
                    "SKU-100005",
                    description_long="An equally long matched-product description from the catalog "
                    "that also needs to wrap instead of colliding with the Qty/UOM/Price columns",
                )],
            ),
            _line(2, raw_description="Short one", selected_item_code=None, candidates=[]),
        ],
    }
    pdf_bytes = render_proposal_pdf(bid)
    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes.startswith(b"%PDF")
    assert len(pdf_bytes) > 0


def test_customer_submission_excel_uses_hs_external_quote_columns():
    """The XLSX mirrors the customer's external customer quote: catalog-named columns
    (Customer Description up front, Manufacturer as its own column), a Date line
    and a Total Bid Price in the header, and a Total row at the bottom."""
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    bid = {
        "bid_id": "bid-test",
        "customer_name": "Test Customer",
        "solicitation_ref": "SOL-1",
        "due_date": "2026-08-01",
        "line_items": [_line(1)],
    }
    workbook = load_workbook(BytesIO(render_customer_submission_excel(bid)))
    sheet = workbook.active

    # Header block: customer, solicitation ref, date.
    assert sheet["A1"].value == "Test Customer"
    assert sheet["A2"].value == "Solicitation Reference: SOL-1"
    assert sheet["A3"].value == "Date: 2026-08-01"

    header_row = [cell.value for cell in sheet[5]][:10]
    assert header_row == [
        "Sort", "Customer Description", "Item #", "Catalog Description",
        "Manufacturer Description", "UOM", "Customer Quantity",
        "Bid Price", "Extended Bid Price", "Comments",
    ]

    data_row = [cell.value for cell in sheet[6]][:10]
    assert data_row == [
        1,
        "Nitrile exam gloves, powder free, size medium, box of 100ct",
        "SKU-100005",
        "Kimberly-Clark Kimtech M nitrile exam gloves, powder-free, textured fingertips, 100/BX",
        "Kimberly-Clark",
        "BX",
        200,
        8.17,
        1634.0,  # 8.17 * 200, rounded
        None,  # Comments blank for a matched line
    ]

    # Total row sums the extended prices; Total Bid Price also appears in the header.
    all_cells = [c.value for row in sheet.iter_rows() for c in row if c.value is not None]
    assert all_cells.count("Total Bid Price:") == 2  # header + total row
    assert all_cells.count(1634.0) == 3  # data extended, header total, total row
